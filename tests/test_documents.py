from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from documents import extract_document, validate_text


def test_utf8_text_is_preserved():
    assert extract_document("needs.TXT", b"  Need 20 seats.\n12 months.  ") == "Need 20 seats.\n12 months."


def test_docx_paragraphs_and_tables_are_extracted():
    doc = Document()
    doc.add_paragraph("Customer requirements")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "20 seats"
    doc.tables[0].rows[0].cells[1].text = "12 months"
    out = BytesIO()
    doc.save(out)
    text = extract_document("rfp.docx", out.getvalue())
    assert "Customer requirements" in text and "20 seats" in text and "12 months" in text


def test_text_pdf_extracts_customer_requirements():
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 700 Td (Need 20 seats for 12 months.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    out = BytesIO()
    writer.write(out)
    assert "Need 20 seats for 12 months." in extract_document("rfp.pdf", out.getvalue())


def test_no_text_pdf_explains_ocr_is_unavailable():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    out = BytesIO()
    writer.write(out)
    with pytest.raises(ValueError, match="text|scanned"):
        extract_document("scan.pdf", out.getvalue())


@pytest.mark.parametrize("name,content", [("a.csv", b"hello"), ("a.txt", b"\xff"),
    ("a.pdf", b"broken"), ("a.docx", b"broken"), ("a.txt", b""),
    ("a.txt", b"x" * (5 * 1024 * 1024 + 1))])
def test_invalid_files_are_rejected(name, content):
    with pytest.raises(ValueError):
        extract_document(name, content)


@pytest.mark.parametrize("text", ["  ", "x" * 40001])
def test_empty_and_oversized_text_are_rejected(text):
    with pytest.raises(ValueError):
        validate_text(text)

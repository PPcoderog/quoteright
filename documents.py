"""Local, bounded text extraction; uploaded files are never persisted."""
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

MAX_BYTES = 5 * 1024 * 1024
MAX_CHARS = 40000


def validate_text(text: str) -> str:
    text = text.strip()
    if not text:
        raise ValueError("No readable text found. Scanned PDFs need OCR; use a text-based document or paste the requirements.")
    if len(text) > MAX_CHARS:
        raise ValueError("The document exceeds 40,000 characters. Use a shorter requirements document.")
    return text


def extract_document(name: str, content: bytes) -> str:
    if len(content) > MAX_BYTES:
        raise ValueError("The file exceeds 5 MB. Upload a smaller document.")
    suffix = Path(name).suffix.lower()
    if suffix not in {".txt", ".pdf", ".docx"}:
        raise ValueError("Upload a PDF, DOCX, or UTF-8 TXT file.")
    try:
        if suffix == ".txt":
            text = content.decode("utf-8-sig")
        elif suffix == ".pdf":
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("This PDF is encrypted. Upload an unlocked copy.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            doc = Document(BytesIO(content))
            text = "\n".join([p.text for p in doc.paragraphs] + [
                " | ".join(cell.text for cell in row.cells)
                for table in doc.tables for row in table.rows])
    except UnicodeDecodeError:
        raise ValueError("The TXT file must use UTF-8 encoding.") from None
    except ValueError:
        raise
    except Exception:
        raise ValueError("Unable to read this document. It may be damaged; try another copy or paste its text.") from None
    return validate_text(text)
